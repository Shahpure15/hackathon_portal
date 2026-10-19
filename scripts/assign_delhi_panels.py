import os
import sys
import argparse
import openpyxl

backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'backend'))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.database import SessionLocal
from app.models import (
    EvaluationRound,
    ProblemStatement,
    PSSelection,
    Team,
    Volunteer,
    User,
    RbacRole,
    UserRoleAssignment,
    JudgeAssignment,
)
from app.auth import get_password_hash

def assign_delhi_panels(apply: bool = False):
    db = SessionLocal()
    mode_str = "APPLY (committing changes to database)" if apply else "DRY RUN (no database changes)"
    print(f"=== Delhi Finale Panel Assignments: {mode_str} ===\n")

    try:
        round_id = "68217b7a-6ad8-4f5f-9079-00c284d275eb"
        round_obj = db.query(EvaluationRound).filter(EvaluationRound.id == round_id).first()
        if not round_obj:
            round_obj = db.query(EvaluationRound).filter(
                (EvaluationRound.roundNumber == 2) | (EvaluationRound.name.ilike('%Delhi%'))
            ).first()

        if not round_obj:
            print(f"Error: Target evaluation round not found.")
            return False

        event_id = round_obj.eventId
        print(f"Target Round: {round_obj.name} (ID: {round_obj.id})")
        print(f"Event ID: {event_id}\n")

        # Load system judge role
        judge_role = db.query(RbacRole).filter(
            RbacRole.eventId == event_id,
            RbacRole.systemKey == "judge"
        ).first()
        if not judge_role:
            print("Warning: System judge role not found. Looking for any role named 'Judge'...")
            judge_role = db.query(RbacRole).filter(
                RbacRole.eventId == event_id,
                RbacRole.name.ilike("Judge")
            ).first()

        # Load problem statements
        ps_list = db.query(ProblemStatement).filter(ProblemStatement.eventId == event_id).all()
        ps_map = {p.code: p for p in ps_list}

        # Load teams
        all_teams = db.query(Team).filter(Team.eventId == event_id).all()
        team_reg_map = {t.registrationId.lower().strip(): t for t in all_teams}
        team_name_map = {t.name.lower().strip(): t for t in all_teams}

        # Load existing volunteers
        existing_vols = db.query(Volunteer).filter(Volunteer.eventId == event_id).all()
        vol_name_map = {v.name.lower().strip(): v for v in existing_vols}

        # Read Excel
        excel_path = os.path.join(os.path.dirname(__file__), '..', 'dump', 'Delhi Data', 'IES_Delhi (1).xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb.active

        excel_rows = []
        excel_judges = set()
        for r in range(2, sheet.max_row + 1):
            ps_no = sheet.cell(r, 2).value
            team_code = sheet.cell(r, 3).value
            team_name = sheet.cell(r, 4).value
            college = sheet.cell(r, 5).value
            rank = sheet.cell(r, 6).value
            m1 = sheet.cell(r, 7).value
            m2 = sheet.cell(r, 8).value
            if team_code:
                j1 = str(m1).strip() if m1 else None
                j2 = str(m2).strip() if m2 else None
                excel_rows.append({
                    "row": r,
                    "psNo": int(ps_no) if ps_no is not None else None,
                    "teamCode": str(team_code).strip(),
                    "teamName": str(team_name).strip() if team_name else "",
                    "judge1": j1,
                    "judge2": j2,
                })
                if j1: excel_judges.add(j1)
                if j2: excel_judges.add(j2)

        print(f"Parsed {len(excel_rows)} team rows and {len(excel_judges)} distinct judges from Excel.")

        # Ensure judges and accounts exist
        judge_vol_map = {}
        default_pwd_hash = get_password_hash("Judge@Delhi2026!")

        for j_name in sorted(list(excel_judges)):
            v = vol_name_map.get(j_name.lower())
            slug = j_name.lower().replace(" ", ".")
            username = f"judge.{slug}"
            email = f"{slug}@delhifinale.local"

            if not v:
                print(f"[Judge Create] Volunteer '{j_name}'")
                v = Volunteer(
                    eventId=event_id,
                    name=j_name,
                    type="JUDGE",
                    contact="9876543210",
                    email=email,
                )
                db.add(v)
                db.flush()
                vol_name_map[j_name.lower()] = v

            judge_vol_map[j_name] = v

            # Check user account
            user = db.query(User).filter(User.eventId == event_id, User.volunteerId == v.id).first()
            if not user:
                user = db.query(User).filter(User.eventId == event_id, User.username == username).first()

            if not user:
                print(f"[User Create] User account '{username}' for {j_name}")
                user = User(
                    eventId=event_id,
                    username=username,
                    passwordHash=default_pwd_hash,
                    role="VOLUNTEER",
                    status="active",
                    volunteerId=v.id,
                )
                db.add(user)
                db.flush()

            # Ensure RBAC role
            if judge_role:
                has_role = db.query(UserRoleAssignment).filter(
                    UserRoleAssignment.eventId == event_id,
                    UserRoleAssignment.userId == user.id,
                    UserRoleAssignment.roleId == judge_role.id
                ).first()
                if not has_role:
                    print(f"[RBAC Assign] Role 'Judge' -> '{username}'")
                    db.add(UserRoleAssignment(
                        eventId=event_id,
                        userId=user.id,
                        roleId=judge_role.id
                    ))

        # Assign teams to judges and verify PS
        assignments_created = 0
        assignments_existing = 0
        ps_assigned = 0

        for item in excel_rows:
            tc = item["teamCode"]
            tn = item["teamName"]
            team = team_reg_map.get(tc.lower()) or team_name_map.get(tn.lower())
            if not team:
                print(f"Error: Team not found for code '{tc}' / name '{tn}'")
                continue

            # Verify PS Selection
            ps_code = f"PSE{item['psNo']:02d}" if item["psNo"] else None
            ps_obj = ps_map.get(ps_code)
            sel = db.query(PSSelection).filter(PSSelection.teamId == team.id).first()

            if not sel and ps_obj:
                print(f"[PS Assign] Team '{team.name}' ({tc}) -> {ps_code}")
                db.add(PSSelection(teamId=team.id, psId=ps_obj.id))
                ps_assigned += 1
            elif sel and ps_obj and sel.psId != ps_obj.id:
                print(f"[PS Update] Team '{team.name}' ({tc}) PS changed to {ps_code}")
                sel.psId = ps_obj.id
                ps_assigned += 1

            # Assign to Judge 1 and Judge 2
            for j_name in (item["judge1"], item["judge2"]):
                if not j_name:
                    continue
                v = judge_vol_map.get(j_name)
                if not v:
                    print(f"Error: Judge volunteer not resolved for '{j_name}'")
                    continue

                existing_assign = db.query(JudgeAssignment).filter(
                    JudgeAssignment.roundId == round_obj.id,
                    JudgeAssignment.teamId == team.id,
                    JudgeAssignment.judgeVolunteerId == v.id,
                ).first()

                if existing_assign:
                    if existing_assign.status != "active":
                        existing_assign.status = "active"
                        existing_assign.revokedAt = None
                        existing_assign.revokedBy = None
                        assignments_created += 1
                    else:
                        assignments_existing += 1
                else:
                    db.add(JudgeAssignment(
                        eventId=event_id,
                        roundId=round_obj.id,
                        teamId=team.id,
                        judgeVolunteerId=v.id,
                        source="wizard",
                        status="active",
                        createdBy="admin_delhi_seed",
                    ))
                    assignments_created += 1

        print(f"\n--- Assignment Summary ---")
        print(f"PS Selections Assigned/Verified: {ps_assigned}")
        print(f"New Judge Assignments: {assignments_created}")
        print(f"Existing Active Assignments: {assignments_existing}")
        print(f"Total Target Assignments: {assignments_created + assignments_existing} (Expected: 58 for 29 teams * 2 judges)")

        if apply:
            db.commit()
            print("\nSUCCESS: All changes committed to the database!")
        else:
            db.rollback()
            print("\nDRY RUN COMPLETE: No changes were committed. Run with '--apply' to commit.")

        return True

    except Exception as e:
        db.rollback()
        print(f"Error executing panel assignments: {e}")
        import traceback
        traceback.print_exc()
        return False
    finally:
        db.close()

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Assign Delhi Finale Panels")
    parser.add_argument("--apply", action="store_true", help="Apply changes to the database")
    args = parser.parse_args()
    assign_delhi_panels(apply=args.apply)
