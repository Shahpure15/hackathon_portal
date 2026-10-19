import os
import sys

backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'backend'))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.database import SessionLocal
from app.models import Volunteer, Team, EvaluationRound, User, Role, UserRoleAssignment, RbacRole
import openpyxl

def check():
    db = SessionLocal()
    round_id = '68217b7a-6ad8-4f5f-9079-00c284d275eb'
    rnd = db.query(EvaluationRound).filter(EvaluationRound.id == round_id).first()

    excel_path = os.path.join(os.path.dirname(__file__), '..', 'dump', 'Delhi Data', 'IES_Delhi (1).xlsx')
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    judges_excel = set()
    teams_excel = []
    for r in range(2, sheet.max_row + 1):
        tc = sheet.cell(r, 3).value
        tn = sheet.cell(r, 4).value
        m1 = sheet.cell(r, 7).value
        m2 = sheet.cell(r, 8).value
        if tc:
            teams_excel.append((str(tc).strip(), str(tn).strip() if tn else ''))
            if m1: judges_excel.add(str(m1).strip())
            if m2: judges_excel.add(str(m2).strip())

    print(f"Distinct Judges in Excel ({len(judges_excel)}):", sorted(list(judges_excel)))
    vols = db.query(Volunteer).filter(Volunteer.eventId == rnd.eventId).all()
    vol_names = {v.name.lower().strip(): v for v in vols}
    for j in sorted(list(judges_excel)):
        v = vol_names.get(j.lower())
        status_str = f"FOUND (id: {v.id}, type: {v.type})" if v else "NOT FOUND"
        print(f"  Judge '{j}': {status_str}")

    db_teams = db.query(Team).filter(Team.eventId == rnd.eventId).all()
    team_reg_map = {t.registrationId.lower().strip(): t for t in db_teams}
    team_name_map = {t.name.lower().strip(): t for t in db_teams}
    print(f"\nTeams in Excel ({len(teams_excel)}):")
    matched = 0
    for tc, tn in teams_excel:
        t = team_reg_map.get(tc.lower()) or team_name_map.get(tn.lower())
        if t:
            matched += 1
        else:
            print(f"  Unmatched: Code '{tc}', Name '{tn}'")
    print(f"Matched teams: {matched}/{len(teams_excel)}")

if __name__ == '__main__':
    check()
