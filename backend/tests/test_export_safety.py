import io
import openpyxl
import pytest
from app.services.export_utils import safe_cell, set_cell_safe, sanitize_worksheet, sanitize_workbook


def test_safe_cell_non_strings():
    assert safe_cell(None) is None
    assert safe_cell(123) == 123
    assert safe_cell(45.67) == 45.67
    assert safe_cell(True) is True
    assert safe_cell(False) is False
    assert safe_cell({}) == {}
    assert safe_cell([]) == []


def test_safe_cell_benign_strings():
    assert safe_cell("") == ""
    assert safe_cell("Normal Text") == "Normal Text"
    assert safe_cell("Team Alpha 101") == "Team Alpha 101"
    assert safe_cell("123-456") == "123-456"


def test_safe_cell_formula_injection():
    # Direct formula prefixes
    assert safe_cell("=1+1") == "'=1+1"
    assert safe_cell("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert safe_cell("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert safe_cell("-2+3") == "'-2+3"
    assert safe_cell("@SUM(1,2)") == "'@SUM(1,2)"
    assert safe_cell("\tTAB_INJECT") == "'\tTAB_INJECT"
    assert safe_cell("\rRETURN_INJECT") == "'\rRETURN_INJECT"

    # Leading whitespace before formula trigger
    assert safe_cell("   =CMD('calc')") == "'   =CMD('calc')"
    assert safe_cell("  +123") == "'  +123"
    assert safe_cell(" \t@HYPERLINK('http://evil.com')") == "' \t@HYPERLINK('http://evil.com')"


def test_set_cell_safe_openpyxl():
    wb = openpyxl.Workbook()
    ws = wb.active

    # Numeric
    c1 = ws.cell(1, 1)
    set_cell_safe(c1, 100)
    assert c1.value == 100
    assert c1.data_type == "n"

    # Plain string
    c2 = ws.cell(1, 2)
    set_cell_safe(c2, "Safe String")
    assert c2.value == "Safe String"
    assert c2.data_type == "s"

    # Formula injection string: value preserved without quote, data_type forced to 's'
    c3 = ws.cell(1, 3)
    set_cell_safe(c3, "=1+1")
    assert c3.value == "=1+1"
    assert c3.data_type == "s"

    c4 = ws.cell(1, 4)
    set_cell_safe(c4, "   @SUM(A1)")
    assert c4.value == "   @SUM(A1)"
    assert c4.data_type == "s"


def test_sanitize_workbook():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws2 = wb.create_sheet(title="Sheet2")

    # In default openpyxl append, =1+1 is classified as formula ('f')
    ws1.append(["=1+1", "+2+2", "plain", 42])
    assert ws1.cell(1, 1).data_type == "f"

    ws2.append(["-3", "@SUM(1,2)", "  =CMD()"])

    # Sanitize workbook
    sanitize_workbook(wb)

    # All formula triggers must now be forced to 's'
    assert ws1.cell(1, 1).data_type == "s"
    assert ws1.cell(1, 1).value == "=1+1"
    assert ws1.cell(1, 2).data_type == "s"
    assert ws1.cell(1, 3).data_type == "s"
    assert ws1.cell(1, 4).data_type == "n"

    assert ws2.cell(1, 1).data_type == "s"
    assert ws2.cell(1, 2).data_type == "s"
    assert ws2.cell(1, 3).data_type == "s"

    # Save to buffer and reload to confirm openpyxl serialization treats it as string
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    reloaded_wb = openpyxl.load_workbook(buf)
    r_ws1 = reloaded_wb["Sheet1"]
    assert r_ws1.cell(1, 1).data_type == "s"
    assert r_ws1.cell(1, 1).value == "=1+1"
