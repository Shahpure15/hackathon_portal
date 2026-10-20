from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import io
import re
import zipfile
from collections import defaultdict
from typing import Dict, List, Optional, Any
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.config import settings
from app.database import get_db
from app.dto.results import (
    ResultsConfigDTO,
    ResultsStatusDTO,
    ResultsUpdateCommitRequest,
    AdvancePreviewRequest,
    AdvanceCommitRequest,
    ProgressionItemDTO,
)
from app.models import (
    EvaluationRound,
    RubricVersion,
    RubricField,
    RubricPSGuidance,
    ResultRelease,
    ResultReleaseItem,
    TeamRoundStatus,
    Team,
    TeamStatus,
    Volunteer,
    ProblemStatement,
    PSSelection,
    JudgeAssignment,
    Evaluation,
    EvaluationVersion,
    EvaluationScore,
    EvaluationText,
)
from app.rbac import Principal, require_permission
from app.routers.admin import log_audit
from app.services.export_utils import safe_cell, set_cell_safe, sanitize_workbook
from app.services.eligibility import is_team_eligible
from app.services.progression import select_top_n, select_threshold
from app.services.ranking import build_leaderboard
from app.services.evaluation import current_finals
from app.services.rubric import applicable_fields
from app.services.results import (
    DEFAULT_RESULTS_CONFIG,
    build_release_items,
    diff_releases,
    release,
    update,
    retract,
)

router = APIRouter(prefix="/admin/rounds/{rid}", tags=["Results & Progression Admin"])


def _get_round(rid: str, event_id: str, db: Session) -> EvaluationRound:
    round_obj = (
        db.query(EvaluationRound)
        .filter(EvaluationRound.id == rid, EvaluationRound.eventId == event_id)
        .first()
    )
    if not round_obj:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Round not found.")
    return round_obj


def _generate_advance_token(rid: str, team_ids: List[str], tie_decision: Optional[str], has_tie: bool = False) -> str:
    sorted_ids = ",".join(sorted(team_ids))
    td = tie_decision or ""
    raw = f"{settings.JWT_SECRET_KEY}:{rid}:{sorted_ids}:{td}:{has_tie}"
    sig = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    return f"{sig}:tie" if has_tie else sig


# ---------------------------------------------------------------------------
# 1. Leaderboard
# ---------------------------------------------------------------------------
@router.get("/leaderboard")
def get_leaderboard(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.view_all")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    return build_leaderboard(db, round_obj)


# ---------------------------------------------------------------------------
# 2. Results Config (GET & PUT)
# ---------------------------------------------------------------------------
@router.get("/results-config", response_model=ResultsConfigDTO)
def get_results_config(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.view_all")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    return round_obj.resultsConfig or DEFAULT_RESULTS_CONFIG


@router.put("/results-config", response_model=ResultsConfigDTO)
def update_results_config(
    rid: str,
    payload: ResultsConfigDTO,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)

    # Validate keys against rubric's text fields
    curr_rubric = (
        db.query(RubricVersion)
        .filter(
            RubricVersion.roundId == rid,
            RubricVersion.version == round_obj.currentRubricVersion,
        )
        .first()
    )
    txt_fields = (
        db.query(RubricField)
        .filter(RubricField.rubricVersionId == curr_rubric.id, RubricField.type == "text")
        .all()
        if curr_rubric
        else []
    )
    valid_keys = {f.fieldKey for f in txt_fields} | {f.id for f in txt_fields}

    for key in (payload.textFields or {}).keys():
        if key not in valid_keys:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unknown text field key in results-config: '{key}'",
            )

    cfg_dict = payload.model_dump()
    round_obj.resultsConfig = cfg_dict
    db.commit()
    db.refresh(round_obj)
    return round_obj.resultsConfig


# ---------------------------------------------------------------------------
# 3. Results Preview & Status
# ---------------------------------------------------------------------------
@router.post("/results/preview")
def preview_results(
    rid: str,
    payload: Optional[ResultsConfigDTO] = None,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    config = payload.model_dump() if payload else (round_obj.resultsConfig or DEFAULT_RESULTS_CONFIG)
    items = build_release_items(db, round_obj, config)

    # Up to 5 sample payloads
    samples = [dict(teamId=tid, **p) for tid, p in list(items.items())[:5]]
    return {
        "samplePayloads": samples,
        "totalTeams": len(items),
    }


@router.get("/results/status", response_model=ResultsStatusDTO)
def get_results_status(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    latest = (
        db.query(ResultRelease)
        .filter(ResultRelease.roundId == rid)
        .order_by(ResultRelease.releaseNo.desc())
        .first()
    )
    if not latest:
        return ResultsStatusDTO(
            releaseNo=0,
            releasedAt=None,
            retracted=False,
            pendingChanges=0,
        )

    retracted = latest.retractedAt is not None
    pending_changes = 0

    if not retracted:
        config = round_obj.resultsConfig or latest.config or DEFAULT_RESULTS_CONFIG
        fresh_items = build_release_items(db, round_obj, config)
        old_items = {it.teamId: it.payload for it in latest.items}
        diff = diff_releases(old_items, fresh_items)
        pending_changes = diff["changedCount"]

    return ResultsStatusDTO(
        releaseNo=latest.releaseNo,
        releasedAt=latest.releasedAt.isoformat() if latest.releasedAt else None,
        retracted=retracted,
        pendingChanges=pending_changes,
    )


# ---------------------------------------------------------------------------
# 4. Results Release, Update Preview, Update, Retract
# ---------------------------------------------------------------------------
@router.post("/results/release")
def release_results(
    rid: str,
    payload: Optional[ResultsConfigDTO] = None,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    config = payload.model_dump() if payload else (round_obj.resultsConfig or DEFAULT_RESULTS_CONFIG)
    if payload:
        round_obj.resultsConfig = config
    rel = release(db, round_obj, config, current_user)
    log_audit(
        db,
        current_user.id,
        "results.released",
        "ResultRelease",
        rel.id,
        {"releaseNo": rel.releaseNo, "totalItems": len(rel.items)},
        event_id=current_user.eventId,
    )
    return {
        "releaseNo": rel.releaseNo,
        "releasedAt": rel.releasedAt.isoformat(),
        "totalItems": len(rel.items),
        "count": len(rel.items),
    }


@router.post("/results/update/preview")
def preview_update_results(
    rid: str,
    payload: Optional[ResultsConfigDTO] = None,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    latest = (
        db.query(ResultRelease)
        .filter(ResultRelease.roundId == rid, ResultRelease.retractedAt == None)
        .order_by(ResultRelease.releaseNo.desc())
        .first()
    )
    if not latest:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No active release found to update.")

    old_items = {it.teamId: it.payload for it in latest.items}
    config = payload.model_dump() if payload else (round_obj.resultsConfig or latest.config or DEFAULT_RESULTS_CONFIG)
    new_items = build_release_items(db, round_obj, config)
    return diff_releases(old_items, new_items)


@router.post("/results/update")
def update_results(
    rid: str,
    payload: ResultsUpdateCommitRequest,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    config = round_obj.resultsConfig or DEFAULT_RESULTS_CONFIG
    rel = update(db, round_obj, config, current_user, payload.expectedReleaseNo)
    log_audit(
        db,
        current_user.id,
        "results.updated",
        "ResultRelease",
        rel.id,
        {"releaseNo": rel.releaseNo},
        event_id=current_user.eventId,
    )
    return {
        "releaseNo": rel.releaseNo,
        "releasedAt": rel.releasedAt.isoformat(),
    }


@router.post("/results/retract")
def retract_results(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.release")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    rel = retract(db, round_obj, current_user)
    log_audit(
        db,
        current_user.id,
        "results.retracted",
        "ResultRelease",
        rel.id,
        {"releaseNo": rel.releaseNo},
        event_id=current_user.eventId,
    )
    return {
        "success": True,
        "retractedAt": rel.retractedAt.isoformat(),
    }


# ---------------------------------------------------------------------------
# 5. Advancement (Preview & Commit)
# ---------------------------------------------------------------------------
@router.post("/advance/preview")
def preview_advance(
    rid: str,
    payload: AdvancePreviewRequest,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.advance")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    lb = build_leaderboard(db, round_obj)
    rows = lb["rows"]

    if payload.mode == "top_n":
        selected_ids, warnings = select_top_n(
            rows,
            payload.n or 0,
            include_provisional=payload.includeProvisional,
            tie_decision=payload.tieDecision,
        )
    elif payload.mode == "threshold":
        selected_ids, warnings = select_threshold(
            rows,
            payload.minScore or Decimal("0"),
            include_provisional=payload.includeProvisional,
        )
    elif payload.mode == "manual":
        selected_ids = payload.teamIds if payload.teamIds is not None else (payload.manualTeamIds or [])
        warnings = []
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported advance mode: {payload.mode}")

    has_tie = any(w.get("type") == "tie_at_cutoff" for w in warnings)
    token = _generate_advance_token(rid, selected_ids, payload.tieDecision, has_tie=has_tie)

    # Build teams display list for the preview UI
    selected_set = set(selected_ids)
    all_teams_map = {t.id: t for t in db.query(Team).filter(Team.eventId == current_user.eventId).all()}
    scored_team_map = {r["teamId"]: r for r in rows}

    teams_preview = []
    for tid in selected_ids:
        t_obj = all_teams_map.get(tid)
        t_row = scored_team_map.get(tid)
        teams_preview.append({
            "id": tid,
            "name": t_obj.name if t_obj else f"Team {tid}",
            "registrationId": t_obj.registrationId if t_obj else "",
            "rank": t_row.get("rank") if t_row else None,
            "score": str(t_row.get("score")) if t_row and t_row.get("score") is not None else None,
            "displayScore": str(t_row.get("displayScore")) if t_row and t_row.get("displayScore") is not None else None,
        })

    return {
        "selectedTeamIds": selected_ids,
        "teams": teams_preview,
        "count": len(selected_ids),
        "hasCutoffTie": has_tie,
        "warnings": warnings,
        "token": token,
    }


@router.post("/advance/commit")
def commit_advance(
    rid: str,
    payload: AdvanceCommitRequest,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.advance")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    lb = build_leaderboard(db, round_obj)
    rows = lb["rows"]
    unscored = lb["unscored"]

    # Validate token
    has_tie = payload.token.endswith(":tie")
    expected_token = _generate_advance_token(rid, payload.teamIds, payload.tieDecision, has_tie=has_tie)
    if payload.token != expected_token:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Advance token mismatch or selection has been tampered with.",
        )

    # Check if a cutoff tie was detected and if tieDecision is required
    if has_tie and payload.tieDecision not in ("include_all_tied", "exclude_tied", "confirmed_manual"):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A tie was detected at the cutoff. You must specify a tieDecision ('include_all_tied', 'exclude_tied', or 'confirmed_manual').",
        )

    selected_set = set(payload.teamIds)
    scored_team_map = {r["teamId"]: r for r in rows}
    all_teams = db.query(Team).filter(Team.eventId == current_user.eventId).all()

    old_statuses = (
        db.query(TeamRoundStatus)
        .filter(TeamRoundStatus.roundId == rid)
        .all()
    )
    old_advanced_count = sum(1 for s in old_statuses if s.state == "advanced")

    existing_status_map = {s.teamId: s for s in old_statuses}
    via_val = f"auto_{payload.mode}" if payload.mode in ("top_n", "threshold") else "manual"
    now_utc = datetime.now(timezone.utc)

    for t in all_teams:
        if t.id in selected_set:
            state = "advanced"
        elif t.id in scored_team_map:
            state = "not_advanced"
        else:
            state = "eligible"

        row_info = scored_team_map.get(t.id)
        # Use custom rank if specified in payload, else fall back to scored rank
        if payload.customRanks and t.id in payload.customRanks:
            rank_snap = payload.customRanks[t.id]
        else:
            rank_snap = row_info["rank"] if row_info else None
        score_snap = row_info["score"] if row_info else None

        existing = existing_status_map.get(t.id)
        if existing:
            existing.state = state
            existing.via = via_val
            existing.rankSnapshot = rank_snap
            existing.scoreSnapshot = score_snap
            existing.decidedBy = current_user.id
            existing.decidedAt = now_utc
        else:
            new_status = TeamRoundStatus(
                eventId=current_user.eventId,
                roundId=rid,
                teamId=t.id,
                state=state,
                via=via_val,
                rankSnapshot=rank_snap,
                scoreSnapshot=score_snap,
                decidedBy=current_user.id,
                decidedAt=now_utc,
            )
            db.add(new_status)

    # If targetRoundId is provided, configure the target round's eligibility to receive these advanced teams
    if payload.targetRoundId:
        target_round = (
            db.query(EvaluationRound)
            .filter(EvaluationRound.id == payload.targetRoundId, EvaluationRound.eventId == current_user.eventId)
            .first()
        )
        if target_round:
            target_round.eligibility = {
                "mode": "advanced_from",
                "fromRoundId": rid,
                "teamIds": list(selected_set),
            }

    log_audit(
        db,
        current_user.id,
        "round.progression_committed",
        "EvaluationRound",
        rid,
        {"oldAdvancedCount": old_advanced_count, "newAdvancedCount": len(selected_set), "targetRoundId": payload.targetRoundId},
        event_id=current_user.eventId,
    )
    db.commit()

    return {
        "success": True,
        "advancedCount": len(selected_set),
        "targetRoundId": payload.targetRoundId,
    }


# ---------------------------------------------------------------------------
# 6. Progression List
# ---------------------------------------------------------------------------
@router.get("/progression", response_model=List[ProgressionItemDTO])
def get_progression(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.view_all")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    statuses = (
        db.query(TeamRoundStatus, Team.name, Team.registrationId)
        .join(Team, TeamRoundStatus.teamId == Team.id)
        .filter(TeamRoundStatus.roundId == rid)
        .order_by(Team.registrationId)
        .all()
    )

    out = []
    for s, t_name, t_reg in statuses:
        out.append(
            ProgressionItemDTO(
                id=s.id,
                teamId=s.teamId,
                name=t_name,
                registrationId=t_reg,
                state=s.state,
                via=s.via,
                rankSnapshot=s.rankSnapshot,
                scoreSnapshot=s.scoreSnapshot,
                decidedBy=s.decidedBy,
                decidedAt=s.decidedAt.isoformat() if s.decidedAt else None,
            )
        )
    return out


# ---------------------------------------------------------------------------
# 7. Leaderboard Export (Formula-hardened, Multi-PS Enriched)
# ---------------------------------------------------------------------------
@router.api_route("/leaderboard/export", methods=["GET", "POST"])
def export_leaderboard(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.export")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    lb = build_leaderboard(db, round_obj)
    rows = lb["rows"]
    unscored = lb["unscored"]

    team_ids = [r["teamId"] for r in rows] + [u["teamId"] for u in unscored]
    teams = db.query(Team).filter(Team.id.in_(team_ids)).all() if team_ids else []
    team_map = {t.id: t for t in teams}

    selections = db.query(PSSelection).filter(PSSelection.teamId.in_(team_ids)).all() if team_ids else []
    team_ps_map = {s.teamId: s.psId for s in selections}

    ps_records = db.query(ProblemStatement).filter(ProblemStatement.eventId == round_obj.eventId).all()
    ps_map = {p.id: p for p in ps_records}

    volunteers = db.query(Volunteer).filter(Volunteer.eventId == round_obj.eventId).all()
    volunteers_map = {v.id: v for v in volunteers}

    assignments = (
        db.query(JudgeAssignment, Volunteer)
        .join(Volunteer, Volunteer.id == JudgeAssignment.judgeVolunteerId)
        .filter(JudgeAssignment.roundId == rid, JudgeAssignment.status == "active")
        .all()
    )
    judges_by_team = defaultdict(list)
    for ja, v in assignments:
        judges_by_team[ja.teamId].append(v.name)

    curr_rubric = (
        db.query(RubricVersion)
        .filter(RubricVersion.roundId == rid, RubricVersion.version == round_obj.currentRubricVersion)
        .first()
    )
    all_rubric_fields = (
        db.query(RubricField)
        .filter(RubricField.rubricVersionId == curr_rubric.id)
        .all()
        if curr_rubric
        else []
    )
    fields_by_id = {f.id: f for f in all_rubric_fields}

    finals = current_finals(db, round_obj)
    finals_by_team = defaultdict(list)
    for f in finals:
        finals_by_team[f.evaluation.teamId].append(f)

    def compute_team_sections(team_id: str):
        t_finals = finals_by_team.get(team_id, [])
        if not t_finals:
            return None, None
        sec_a_totals = []
        sec_b_totals = []
        for fin in t_finals:
            a_sum = 0.0
            b_sum = 0.0
            for sc in fin.scores:
                fld = fields_by_id.get(sc.fieldId)
                if not fld or fld.type != "numeric":
                    continue
                m = float(sc.marks or 0)
                if fld.problemStatementId:
                    a_sum += m
                else:
                    b_sum += m
            sec_a_totals.append(a_sum)
            sec_b_totals.append(b_sum)
        avg_a = round(sum(sec_a_totals) / len(sec_a_totals), 2) if sec_a_totals else 0.0
        avg_b = round(sum(sec_b_totals) / len(sec_b_totals), 2) if sec_b_totals else 0.0
        return avg_a, avg_b

    def get_panel_str(team_id: str):
        count = len(judges_by_team.get(team_id, []))
        for fin in finals_by_team.get(team_id, []):
            vid = getattr(fin.evaluation, 'judgeVolunteerId', None)
            if vid and vid in volunteers_map and vid not in judges_by_team.get(team_id, []):
                count += 1
        return ", ".join([f"Judge {i+1}" for i in range(count)]) or "Unassigned"

    has_shortlisted_flag = any(
        (t.status == TeamStatus.SHORTLISTED or str(t.status).upper() == "SHORTLISTED")
        for t in teams
    )
    def is_team_included(t_id: str) -> bool:
        t_obj = team_map.get(t_id)
        if not t_obj:
            return False
        if has_shortlisted_flag:
            is_st = (t_obj.status == TeamStatus.SHORTLISTED or str(t_obj.status).upper() == "SHORTLISTED")
            if not is_st:
                return False
        else:
            st_str = str(t_obj.status).upper()
            if "DISQUALIFIED" in st_str or "WAITLISTED" in st_str:
                return False
        return is_team_eligible(db, round_obj, t_id)

    rows = [r for r in rows if is_team_included(r["teamId"])]
    unscored = [u for u in unscored if is_team_included(u["teamId"])]

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws = wb.create_sheet(title=f"Round {round_obj.roundNumber} Leaderboard"[:31])

    headers = [
        "Rank",
        "Team Registration ID",
        "Team Name",
        "College / Institution",
        "Team Leader",
        "Leader Phone",
        "Problem Statement Code",
        "Problem Statement Title",
        "Panel Judges",
        "Section A (PS Specific)",
        "Section B (General Criteria)",
        "Final Score (/100)",
        "Judges Evaluated",
        "Evaluation Status",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 2
    for r in rows:
        t_obj = team_map.get(r["teamId"])
        ps_id = team_ps_map.get(r["teamId"])
        ps_obj = ps_map.get(ps_id) if ps_id else None
        ps_code = ps_obj.code if ps_obj else "-"
        ps_title = ps_obj.title if ps_obj else "Not Assigned"
        panel_str = get_panel_str(r["teamId"])
        sec_a, sec_b = compute_team_sections(r["teamId"])
        status_label = "Provisional" if r.get("provisional") else "Finalized"

        ws.append([
            r["rank"],
            r["registrationId"],
            r["name"],
            t_obj.collegeName if t_obj else "-",
            t_obj.leaderName if t_obj else "-",
            t_obj.leaderPhone if t_obj else "-",
            ps_code,
            ps_title,
            panel_str,
            sec_a if sec_a is not None else float(r["score"]),
            sec_b if sec_b is not None else 0.0,
            float(r["displayScore"]) if r.get("displayScore") is not None else float(r["score"]),
            r["judgesCount"],
            status_label,
        ])
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 10, 11, 12, 13, 14):
                cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

    for u in unscored:
        t_obj = team_map.get(u["teamId"])
        ps_id = team_ps_map.get(u["teamId"])
        ps_obj = ps_map.get(ps_id) if ps_id else None
        ps_code = ps_obj.code if ps_obj else "-"
        ps_title = ps_obj.title if ps_obj else "Not Assigned"
        panel_str = get_panel_str(u["teamId"])

        ws.append([
            "-",
            u["registrationId"],
            u["name"],
            t_obj.collegeName if t_obj else "-",
            t_obj.leaderName if t_obj else "-",
            t_obj.leaderPhone if t_obj else "-",
            ps_code,
            ps_title,
            panel_str,
            "-",
            "-",
            "-",
            0,
            "Unscored",
        ])
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 10, 11, 12, 13, 14):
                cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    sanitize_workbook(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"round_{round_obj.roundNumber}_leaderboard.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



# ---------------------------------------------------------------------------
# 8. Print-Ready Multi-Sheet Detailed Workbook Export
# ---------------------------------------------------------------------------
@router.api_route("/export/detailed-sheets", methods=["GET", "POST"])
def export_detailed_sheets(
    rid: str,
    db: Session = Depends(get_db),
    current_user: Principal = Depends(require_permission("evaluation.export")),
):
    round_obj = _get_round(rid, current_user.eventId, db)
    lb = build_leaderboard(db, round_obj)
    rows = lb["rows"]
    unscored = lb["unscored"]
    all_entries = rows + unscored

    curr_rubric = (
        db.query(RubricVersion)
        .filter(RubricVersion.roundId == rid, RubricVersion.version == round_obj.currentRubricVersion)
        .first()
    )
    if not curr_rubric:
        raise HTTPException(status_code=400, detail="No rubric published for this round.")

    team_ids = [e["teamId"] for e in all_entries]
    teams = db.query(Team).filter(Team.id.in_(team_ids)).all() if team_ids else []
    team_map = {t.id: t for t in teams}

    # Filter strictly to teams shortlisted and eligible for this round
    has_shortlisted_flag = any(
        (t.status == TeamStatus.SHORTLISTED or str(t.status).upper() == "SHORTLISTED")
        for t in teams
    )
    shortlisted_entries = []
    for e in all_entries:
        t_id = e["teamId"]
        t_obj = team_map.get(t_id)
        if not t_obj:
            continue
        if has_shortlisted_flag:
            is_st = (t_obj.status == TeamStatus.SHORTLISTED or str(t_obj.status).upper() == "SHORTLISTED")
            if not is_st:
                continue
        else:
            st_str = str(t_obj.status).upper()
            if "DISQUALIFIED" in st_str or "WAITLISTED" in st_str:
                continue
        if not is_team_eligible(db, round_obj, t_id):
            continue
        shortlisted_entries.append(e)

    ps_records = db.query(ProblemStatement).filter(ProblemStatement.eventId == round_obj.eventId).all()
    ps_map = {p.id: p for p in ps_records}

    selections = db.query(PSSelection).filter(PSSelection.teamId.in_(team_ids)).all() if team_ids else []
    team_ps_map = {s.teamId: s.psId for s in selections}

    all_rubric_fields = db.query(RubricField).filter(RubricField.rubricVersionId == curr_rubric.id).all()
    fields_by_id = {f.id: f for f in all_rubric_fields}

    guidance_records = db.query(RubricPSGuidance).filter(RubricPSGuidance.rubricVersionId == curr_rubric.id).all()
    guidance_by_ps = {g.problemStatementId: g for g in guidance_records}

    volunteers = db.query(Volunteer).filter(Volunteer.eventId == round_obj.eventId).all()
    volunteers_map = {v.id: v for v in volunteers}

    assignments = (
        db.query(JudgeAssignment, Volunteer)
        .join(Volunteer, Volunteer.id == JudgeAssignment.judgeVolunteerId)
        .filter(JudgeAssignment.roundId == rid, JudgeAssignment.status == "active")
        .all()
    )
    judges_by_team = defaultdict(list)
    for ja, v in assignments:
        judges_by_team[ja.teamId].append(v)

    finals = current_finals(db, round_obj)
    finals_by_team = defaultdict(list)
    for f in finals:
        finals_by_team[f.evaluation.teamId].append(f)

    # Styles
    navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    dark_gray_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    slate_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    emerald_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_banner = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_data = Font(name="Calibri", size=9.5, color="1E293B")
    font_grand = Font(name="Calibri", size=11, bold=True, color="166534")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    thick_bottom_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="double", color="166534"),
    )

    # Prepare ZIP Archive containing Leaderboard.xlsx and individual team excels
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # -------------------------------------------------------------
        # 1. Leaderboard & Evaluation Summary Workbook
        # -------------------------------------------------------------
        wb_sum = openpyxl.Workbook()
        ws_sum = wb_sum.active
        ws_sum.title = "Evaluation Summary"

        ws_sum.merge_cells("A1:L1")
        top_sum = ws_sum.cell(row=1, column=1, value=f"{round_obj.name.upper()} - EVALUATION SUMMARY & LEADERBOARD")
        top_sum.font = font_banner
        top_sum.fill = navy_fill
        top_sum.alignment = Alignment(horizontal="center", vertical="center")
        ws_sum.row_dimensions[1].height = 30

        sum_headers = [
            "Rank", "Team Code", "Team Name", "College / Institution", "Team Leader",
            "PS Code", "Problem Statement Title", "Assigned Panel Judges",
            "Sec A Score", "Sec B Score", "Total Score (/100)", "Status"
        ]
        ws_sum.append([])
        ws_sum.append(sum_headers)
        for col_idx in range(1, len(sum_headers) + 1):
            c = ws_sum.cell(row=3, column=col_idx)
            c.font = font_header
            c.fill = dark_gray_fill
            c.alignment = Alignment(horizontal="center", vertical="center")

        curr_sum_row = 4
        for e in shortlisted_entries:
            t_id = e["teamId"]
            t_obj = team_map.get(t_id)
            ps_id = team_ps_map.get(t_id)
            ps_obj = ps_map.get(ps_id) if ps_id else None
            ps_code = ps_obj.code if ps_obj else "-"
            ps_title = ps_obj.title if ps_obj else "Not Assigned"

            # Anonymise panel judges in summary
            p_count = len(judges_by_team.get(t_id, []))
            for fin in finals_by_team.get(t_id, []):
                vid = getattr(fin.evaluation, 'judgeVolunteerId', None)
                if vid and vid in volunteers_map and vid not in [j.id for j in judges_by_team.get(t_id, [])]:
                    p_count += 1
            panel_str = ", ".join([f"Judge {i+1}" for i in range(p_count)]) or "Unassigned"

            # Calculate section marks
            t_finals = finals_by_team.get(t_id, [])
            sec_a_avg = None
            sec_b_avg = None
            if t_finals:
                a_sums = []
                b_sums = []
                for fin in t_finals:
                    a_m = sum(float(sc.marks or 0) for sc in fin.scores if sc.marks is not None and fields_by_id.get(sc.fieldId) and fields_by_id[sc.fieldId].problemStatementId)
                    b_m = sum(float(sc.marks or 0) for sc in fin.scores if sc.marks is not None and fields_by_id.get(sc.fieldId) and not fields_by_id[sc.fieldId].problemStatementId and fields_by_id[sc.fieldId].type == 'numeric')
                    a_sums.append(a_m)
                    b_sums.append(b_m)
                sec_a_avg = round(sum(a_sums) / len(a_sums), 2) if a_sums else 0.0
                sec_b_avg = round(sum(b_sums) / len(b_sums), 2) if b_sums else 0.0

            rank_val = e.get("rank") or "-"
            final_score_val = float(e["displayScore"]) if e.get("displayScore") is not None else (float(e["score"]) if e.get("score") is not None else "-")
            status_val = "Finalized" if not e.get("provisional") and e.get("score") is not None else ("Provisional" if e.get("score") is not None else "Unscored")

            ws_sum.append([
                rank_val,
                e["registrationId"],
                e["name"],
                t_obj.collegeName if t_obj else "-",
                t_obj.leaderName if t_obj else "-",
                ps_code,
                ps_title,
                panel_str,
                sec_a_avg if sec_a_avg is not None else "-",
                sec_b_avg if sec_b_avg is not None else "-",
                final_score_val,
                status_val,
            ])
            for col_idx in range(1, len(sum_headers) + 1):
                c = ws_sum.cell(row=curr_sum_row, column=col_idx)
                c.font = font_data
                c.border = thin_border
                if curr_sum_row % 2 == 1:
                    c.fill = zebra_fill
                if col_idx in (1, 6, 9, 10, 11, 12):
                    c.alignment = Alignment(horizontal="center", vertical="center")
            curr_sum_row += 1

        for col in ws_sum.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sum.column_dimensions[col_letter].width = max(max_len + 3, 12)

        sanitize_workbook(wb_sum)
        sum_buf = io.BytesIO()
        wb_sum.save(sum_buf)
        zip_file.writestr(f"Round_{round_obj.roundNumber}_Evaluation/Leaderboard.xlsx", sum_buf.getvalue())

        # -------------------------------------------------------------
        # 2. Individual Team Workbooks (Separate Files, Anonymised)
        # -------------------------------------------------------------
        for e in shortlisted_entries:
            t_id = e["teamId"]
            t_obj = team_map.get(t_id)
            if not t_obj:
                continue

            wb_team = openpyxl.Workbook()
            ws_team = wb_team.active
            ws_team.title = "Evaluation Sheet"

            # Page setup (Portrait, A4, fit to 1 page wide)
            ws_team.page_setup.orientation = ws_team.ORIENTATION_PORTRAIT
            ws_team.page_setup.paperSize = ws_team.PAPERSIZE_A4
            ws_team.page_setup.fitToPage = True
            ws_team.page_setup.fitToWidth = 1
            ws_team.page_setup.fitToHeight = 0
            ws_team.sheet_properties.pageSetUpPr.fitToPage = True

            # Resolve judges for this team
            t_finals = finals_by_team.get(t_id, [])
            eval_by_vol_id = {}
            for fin in t_finals:
                if hasattr(fin, 'evaluation') and fin.evaluation and fin.evaluation.judgeVolunteerId:
                    eval_by_vol_id[fin.evaluation.judgeVolunteerId] = fin

            combined_judges = []
            seen_judge_ids = set()
            for fin in t_finals:
                vid = getattr(fin.evaluation, 'judgeVolunteerId', None)
                if vid and vid not in seen_judge_ids and vid in volunteers_map:
                    combined_judges.append(volunteers_map[vid])
                    seen_judge_ids.add(vid)

            for j in judges_by_team.get(t_id, []):
                if j.id not in seen_judge_ids:
                    combined_judges.append(j)
                    seen_judge_ids.add(j.id)

            num_judges = max(len(combined_judges), 2)
            # ANONYMISED JUDGE NAMES: Judge 1, Judge 2, etc. (No real names)
            judge_names = [f"Judge {i+1}" for i in range(num_judges)]

            total_cols = 5 + num_judges + 1
            last_col_letter = get_column_letter(total_cols)

            # Header Banner
            ws_team.merge_cells(f"A1:{last_col_letter}1")
            top_c = ws_team.cell(row=1, column=1, value=f"{round_obj.name.upper()} - OFFICIAL EVALUATION SHEET")
            top_c.font = font_banner
            top_c.fill = navy_fill
            top_c.alignment = Alignment(horizontal="center", vertical="center")
            ws_team.row_dimensions[1].height = 28

            # Team Info Block
            ps_id = team_ps_map.get(t_id)
            ps_obj = ps_map.get(ps_id) if ps_id else None
            ps_code = ps_obj.code if ps_obj else "PS"
            ps_title = ps_obj.title if ps_obj else "Not Assigned"

            mid_col = max(3, total_cols // 2)
            mid_col_letter = get_column_letter(mid_col)

            ws_team.cell(row=3, column=1, value="Team Name:").font = font_bold
            ws_team.merge_cells(f"B3:{mid_col_letter}3")
            ws_team.cell(row=3, column=2, value=f"{t_obj.name} ({t_obj.collegeName})").font = font_data

            ws_team.cell(row=3, column=mid_col + 1, value="Registration ID:").font = font_bold
            ws_team.merge_cells(f"{get_column_letter(mid_col + 2)}3:{last_col_letter}3")
            ws_team.cell(row=3, column=mid_col + 2, value=t_obj.registrationId or "-").font = font_data

            ws_team.cell(row=4, column=1, value="Problem Statement:").font = font_bold
            ws_team.merge_cells(f"B4:{mid_col_letter}4")
            ws_team.cell(row=4, column=2, value=f"[{ps_code}] {ps_title}").font = font_data

            ws_team.cell(row=4, column=mid_col + 1, value="Panel Judges:").font = font_bold
            ws_team.merge_cells(f"{get_column_letter(mid_col + 2)}4:{last_col_letter}4")
            ws_team.cell(row=4, column=mid_col + 2, value=", ".join(judge_names) or "Unassigned").font = font_data

            # Guidance / Decisive Test Banner
            guidance_obj = guidance_by_ps.get(ps_id) if ps_id else None
            curr_row = 6
            if guidance_obj and (guidance_obj.decisiveTest or guidance_obj.notes):
                ws_team.cell(row=curr_row, column=1, value=f"DECISIVE STRESS TEST / REQUIRED CASE ({ps_code}):").font = Font(name="Calibri", size=9.5, bold=True, color="B45309")
                curr_row += 1
                ws_team.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=total_cols)
                dec_text = guidance_obj.decisiveTest or ""
                if guidance_obj.notes:
                    dec_text += f"\nJudging Guidance Notes: {guidance_obj.notes}"
                c_dec = ws_team.cell(row=curr_row, column=1, value=dec_text)
                c_dec.font = Font(name="Calibri", size=9, italic=True, color="78350F")
                c_dec.fill = amber_fill
                c_dec.alignment = Alignment(wrap_text=True, vertical="center")
                c_dec.border = thin_border
                ws_team.row_dimensions[curr_row].height = 40
                curr_row += 2

            # Criteria Breakdown Table
            app_fields = applicable_fields(db, curr_rubric, t_id, require_ps=False)
            sec_a_fields = [f for f in app_fields if f.type == "numeric" and f.problemStatementId]
            sec_b_fields = [f for f in app_fields if f.type == "numeric" and not f.problemStatementId]
            # Exclude internal staff notes! Only export participant-visible text comments
            txt_fields = [f for f in app_fields if f.type == "text" and f.visibility == "participant"]

            judge_scores_map = []
            judge_texts_map = []
            for i in range(num_judges):
                v_id = combined_judges[i].id if i < len(combined_judges) else None
                ev = eval_by_vol_id.get(v_id) if v_id else None
                if ev:
                    sc_dict = {sc.fieldId: float(sc.marks) for sc in ev.scores if sc.marks is not None}
                    tx_dict = {tx.fieldId: tx.value for tx in ev.texts if tx.value is not None}
                else:
                    sc_dict = {}
                    tx_dict = {}
                judge_scores_map.append(sc_dict)
                judge_texts_map.append(tx_dict)

            table_headers = ["Section", "Code", "Evaluation Criterion", "Max", "Weight"]
            for i in range(num_judges):
                table_headers.append(f"J{i+1}")
            table_headers.append("Avg Marks")

            ws_team.row_dimensions[curr_row].height = 22
            for col_i, h in enumerate(table_headers, 1):
                c = ws_team.cell(row=curr_row, column=col_i, value=h)
                c.font = font_header
                c.fill = dark_gray_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = thin_border
            curr_row += 1

            # Section A rows
            j_a_subtotals = [0.0] * num_judges
            avg_a_subtotal = 0.0

            if sec_a_fields:
                for idx_a, f in enumerate(sec_a_fields, 1):
                    row_scores = [judge_scores_map[i].get(f.id) for i in range(num_judges)]
                    valid_vals = [v for v in row_scores if v is not None]
                    avg_m = round(sum(valid_vals) / len(valid_vals), 2) if valid_vals else None

                    for i, v in enumerate(row_scores):
                        if v is not None:
                            j_a_subtotals[i] += v
                    if avg_m is not None:
                        avg_a_subtotal += avg_m

                    row_vals = [
                        f"Sec A ({ps_code})",
                        f"A.{idx_a}",
                        f.label,
                        f.maxMarks,
                        f"{f.weight}%",
                    ]
                    for v in row_scores:
                        row_vals.append(v if v is not None else "-")
                    row_vals.append(avg_m if avg_m is not None else "-")

                    for col_i, rv in enumerate(row_vals, 1):
                        c = ws_team.cell(row=curr_row, column=col_i, value=rv)
                        c.font = font_data
                        c.border = thin_border
                        c.alignment = Alignment(horizontal="center" if col_i not in (1, 3) else "left", vertical="center")
                    curr_row += 1

                # Subtotal Section A Row
                sec_a_max = sum(float(f.maxMarks or 0) for f in sec_a_fields)
                sec_a_wt = sum(float(f.weight or 0) for f in sec_a_fields)

                ws_team.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
                sub_a_label = ws_team.cell(row=curr_row, column=1, value=f"SECTION A: PROBLEM STATEMENT SPECIFIC SUBTOTAL (MAX {int(sec_a_max)})")
                sub_a_label.font = font_bold
                sub_a_label.alignment = Alignment(horizontal="right", vertical="center")
                ws_team.cell(row=curr_row, column=4, value=sec_a_max).font = font_bold
                ws_team.cell(row=curr_row, column=5, value=f"{int(sec_a_wt)}%").font = font_bold

                for i in range(num_judges):
                    has_eval = any(judge_scores_map[i].get(f.id) is not None for f in sec_a_fields)
                    ws_team.cell(row=curr_row, column=6 + i, value=round(j_a_subtotals[i], 2) if has_eval else "-").font = font_bold

                ws_team.cell(row=curr_row, column=total_cols, value=round(avg_a_subtotal, 2) if t_finals else "-").font = font_bold

                for col_i in range(1, total_cols + 1):
                    c = ws_team.cell(row=curr_row, column=col_i)
                    c.fill = amber_fill
                    c.border = thin_border
                    if col_i in range(4, total_cols + 1):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                curr_row += 1
            else:
                sec_a_max = 0.0
                sec_a_wt = 0.0

            # Section B rows
            j_b_subtotals = [0.0] * num_judges
            avg_b_subtotal = 0.0

            for idx_b, f in enumerate(sec_b_fields, 1):
                row_scores = [judge_scores_map[i].get(f.id) for i in range(num_judges)]
                valid_vals = [v for v in row_scores if v is not None]
                avg_m = round(sum(valid_vals) / len(valid_vals), 2) if valid_vals else None

                for i, v in enumerate(row_scores):
                    if v is not None:
                        j_b_subtotals[i] += v
                if avg_m is not None:
                    avg_b_subtotal += avg_m

                row_vals = [
                    "Sec B (General)",
                    f"B.{idx_b}",
                    f.label,
                    f.maxMarks,
                    f"{f.weight}%",
                ]
                for v in row_scores:
                    row_vals.append(v if v is not None else "-")
                row_vals.append(avg_m if avg_m is not None else "-")

                for col_i, rv in enumerate(row_vals, 1):
                    c = ws_team.cell(row=curr_row, column=col_i, value=rv)
                    c.font = font_data
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="center" if col_i not in (1, 3) else "left", vertical="center")
                curr_row += 1

            # Subtotal Section B Row
            sec_b_max = sum(float(f.maxMarks or 0) for f in sec_b_fields)
            sec_b_wt = sum(float(f.weight or 0) for f in sec_b_fields)

            ws_team.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
            sub_b_label = ws_team.cell(row=curr_row, column=1, value=f"SECTION B: GENERAL CRITERIA SUBTOTAL (MAX {int(sec_b_max)})")
            sub_b_label.font = font_bold
            sub_b_label.alignment = Alignment(horizontal="right", vertical="center")
            ws_team.cell(row=curr_row, column=4, value=sec_b_max).font = font_bold
            ws_team.cell(row=curr_row, column=5, value=f"{int(sec_b_wt)}%").font = font_bold

            for i in range(num_judges):
                has_eval = any(judge_scores_map[i].get(f.id) is not None for f in sec_b_fields)
                ws_team.cell(row=curr_row, column=6 + i, value=round(j_b_subtotals[i], 2) if has_eval else "-").font = font_bold

            ws_team.cell(row=curr_row, column=total_cols, value=round(avg_b_subtotal, 2) if t_finals else "-").font = font_bold

            for col_i in range(1, total_cols + 1):
                c = ws_team.cell(row=curr_row, column=col_i)
                c.fill = slate_fill
                c.border = thin_border
                if col_i in range(4, total_cols + 1):
                    c.alignment = Alignment(horizontal="center", vertical="center")
            curr_row += 1

            # Grand Total Score Row
            grand_max = int(sec_a_max + sec_b_max)
            grand_wt = int(sec_a_wt + sec_b_wt)

            ws_team.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
            total_label = ws_team.cell(row=curr_row, column=1, value=f"FINAL TOTAL EVALUATION SCORE (MAX {grand_max})")
            total_label.font = font_grand
            total_label.alignment = Alignment(horizontal="right", vertical="center")
            ws_team.cell(row=curr_row, column=4, value=grand_max).font = font_grand
            ws_team.cell(row=curr_row, column=5, value=f"{grand_wt}%").font = font_grand

            for i in range(num_judges):
                v_id = combined_judges[i].id if i < len(combined_judges) else None
                has_eval = v_id in eval_by_vol_id if v_id else False
                j_tot = round(j_a_subtotals[i] + j_b_subtotals[i], 2)
                ws_team.cell(row=curr_row, column=6 + i, value=j_tot if has_eval else "-").font = font_grand

            final_team_score = round(avg_a_subtotal + avg_b_subtotal, 2)
            ws_team.cell(row=curr_row, column=total_cols, value=final_team_score if t_finals else "-").font = font_grand

            for col_i in range(1, total_cols + 1):
                c = ws_team.cell(row=curr_row, column=col_i)
                c.fill = emerald_fill
                c.border = thick_bottom_border
                if col_i in range(4, total_cols + 1):
                    c.alignment = Alignment(horizontal="center", vertical="center")
            curr_row += 2

            # Participant Visible Remarks Only (No Internal Staff Notes!)
            if txt_fields:
                ws_team.cell(row=curr_row, column=1, value="JUDGE QUALITATIVE REMARKS & WRITTEN OBSERVATIONS:").font = font_bold
                curr_row += 1

                for tf in txt_fields:
                    ws_team.cell(row=curr_row, column=1, value=f"Criterion: {tf.label}").font = font_bold
                    curr_row += 1

                    for i in range(num_judges):
                        j_name = judge_names[i]
                        val = judge_texts_map[i].get(tf.id) or "None recorded."
                        ws_team.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=total_cols)
                        c_txt = ws_team.cell(row=curr_row, column=1, value=f"{j_name}: {val}")
                        c_txt.font = font_data
                        c_txt.fill = zebra_fill
                        c_txt.border = thin_border
                        c_txt.alignment = Alignment(wrap_text=True, vertical="top")
                        ws_team.row_dimensions[curr_row].height = 28
                        curr_row += 1
                    curr_row += 1

            # Signature & Sign-off Block (Anonymised)
            curr_row += 1
            ws_team.cell(row=curr_row, column=1, value="OFFICIAL VERIFICATION & SIGN-OFF:").font = font_bold
            curr_row += 1

            col_per_judge = max(3, total_cols // num_judges)
            for i in range(num_judges):
                start_c = 1 + i * col_per_judge
                end_c = min(total_cols, start_c + col_per_judge - 1)
                if i == num_judges - 1:
                    end_c = total_cols

                ws_team.merge_cells(start_row=curr_row, start_column=start_c, end_row=curr_row, end_column=end_c)
                sig_cell = ws_team.cell(
                    row=curr_row,
                    column=start_c,
                    value=f"Panel Judge {i+1}\nSignature: ________________________________\nDate: ________________________",
                )
                sig_cell.font = font_data
                sig_cell.alignment = Alignment(wrap_text=True, vertical="top")
                sig_cell.border = thin_border

            ws_team.row_dimensions[curr_row].height = 54

            # Column widths
            ws_team.column_dimensions["A"].width = 16
            ws_team.column_dimensions["B"].width = 8
            ws_team.column_dimensions["C"].width = 38
            ws_team.column_dimensions["D"].width = 10
            ws_team.column_dimensions["E"].width = 10
            for i in range(num_judges):
                col_letter = get_column_letter(6 + i)
                ws_team.column_dimensions[col_letter].width = 18
            ws_team.column_dimensions[last_col_letter].width = 16

            sanitize_workbook(wb_team)
            team_buf = io.BytesIO()
            wb_team.save(team_buf)
            safe_reg = re.sub(r'[\\/*?:\[\]<>|"]', '_', t_obj.registrationId or 'REG')
            safe_name = re.sub(r'[\\/*?:\[\]<>|"]', '_', t_obj.name)[:30]
            rank_str = str(e.get('rank') or 'NA')
            team_filename = f"Rank_{rank_str}_{safe_reg}_{safe_name}.xlsx"
            zip_file.writestr(f"Round_{round_obj.roundNumber}_Evaluation/Teams/{team_filename}", team_buf.getvalue())

    zip_buf.seek(0)
    filename = f"round_{round_obj.roundNumber}_evaluation_sheets.zip"
    return StreamingResponse(
        zip_buf,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


